# encoding: utf-8
import base64
from io import BytesIO
from fontTools.ttLib import TTFont
import requests
import re
from lxml import etree
import re,random
import pandas as pd
from openpyxl.workbook import Workbook
from xpinyin import Pinyin
import time

def get_page(url):
#url = 'https://zk.58.com/chuzu/pn3/?utm_source=market&spm=b-31580022738699-me-f-862.mingzhan&PGTID=0d3090a7-003a-5338-3ae4-adbbf66a341c&ClickID=2'
    headers = {
        "accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3",
        "accept-encoding": "gzip, deflate, br",
        "accept-language": "zh-CN,zh;q=0.9",
        "cache-control": "max-age=0",
        "upgrade-insecure-requests": "1",
        "user-agent": "Mozilla/5.0 (Windows NT 6.1; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/73.0.3683.86 Safari/537.36",
    }

    res = requests.get(url, headers=headers)
    #bs64_str = re.findall("charset=utf-8;base64,(.*?)'\)", res.text)[0]
    return res


def get_page_show_ret(res,string):
    bs64_str = re.findall("charset=utf-8;base64,(.*?)'\)", res.text)[0]
    font = TTFont(BytesIO(base64.decodebytes(bs64_str.encode())))
    c = font['cmap'].tables[0].ttFont.tables['cmap'].tables[0].cmap
    ret_list = []
    for char in string:
        decode_num = ord(char)
        if decode_num in c:
            num = c[decode_num]
            num = int(num[-2:]) - 1
            ret_list.append(num)
        else:
            ret_list.append(char)
    ret_str_show = ''
    for num in ret_list:
        ret_str_show += str(num)
    return ret_str_show

def getinfo_url(res):

    page = etree.HTML(res.text)
    li = page.xpath('.//ul[@class="house-list"]//li')[0:-1]
    url_list = []
    for each_li in li:
        # title1 = each_li.xpath('.//div[@class="des"]/h2/a/text()')[0].strip()
        # title = get_page_show_ret(res,title1)
        # price1 = each_li.xpath('.//div[@class="money"]/b/text()')[0]
        # price = get_page_show_ret(res,price1)
        info_url = each_li.xpath('.//div[@class="des"]/h2/a')[0].get('href')
        # print(info_url)
        # print('='*50)
        url_list.append(info_url)
    return url_list

def getinfo(url_list):
    getinfo_page = []

    for url in url_list:
        gfp = {}
        print(url)
        res = get_page(url)
        page = etree.HTML(res.text)
        try:
            if len(page.xpath('.//div[@class="house-title"]/h1/text()')):
                gfp['标题'] = get_page_show_ret(res, page.xpath('.//div[@class="house-title"]/h1/text()')[0])
            else:
                gfp['标题'] = ' '

            if len(page.xpath('.//ul[@class="f14"]/li/span[2]/text()'))!=0:
                gfp['租赁方式'] = get_page_show_ret(res, page.xpath('.//ul[@class="f14"]/li/span[2]/text()')[0])
            else:
                gfp['租赁方式'] = ' '
            if len(page.xpath('.//ul[@class="f14"]/li[2]/span[2]/text()'))!=0:
                gfp['房屋类型'] = get_page_show_ret(res, page.xpath('.//ul[@class="f14"]/li[2]/span[2]/text()')[0]).strip().replace(
                    ' ', '').replace('\xa0\xa0', '')
            else:
                gfp['房屋类型'] = ' '

            if len(page.xpath('.//ul[@class="f14"]/li[3]/span[2]/text()'))!=0:
                gfp['朝向楼层'] = get_page_show_ret(res, page.xpath('.//ul[@class="f14"]/li[3]/span[2]/text()')[0].strip()).replace(
                    '\xa0\xa0', ' ')
            else:
                gfp['朝向楼层'] = ' '

            if len(page.xpath('.//ul[@class="f14"]/li[4]//a/text()'))!=0:
                gfp['所在小区'] = get_page_show_ret(res, page.xpath('.//ul[@class="f14"]/li[4]//a/text()')[0].strip())
            else:
                gfp['所在小区'] = ' '
            if len(page.xpath('.//ul[@class="f14"]/li[6]/span[2]/text()'))!=0:
                gfp['详细地址'] = get_page_show_ret(res, page.xpath('.//ul[@class="f14"]/li[6]/span[2]/text()')[0].strip())
            else:
                gfp['详细地址'] = ' '

            if len(page.xpath('.//div[@class="house-chat-phone"]/span/text()'))!=0:
                gfp['联系电话'] = get_page_show_ret(res, page.xpath('.//div[@class="house-chat-phone"]/span/text()')[0].strip())
            else:
                gfp['联系电话'] = ' '


            if len(page.xpath('.//div[@class="house-pay-way f16"]/span[2]/text()'))!=0:
                gfp['付款方式'] = get_page_show_ret(res,page.xpath('.//div[@class="house-pay-way f16"]/span[2]/text()')[0])
            else:
                gfp['付款方式'] =' '

            if len(page.xpath('.//div[@class="house-pay-way f16"]//text()'))!=0:
                gfp['价格'] = get_page_show_ret(res, page.xpath('.//div[@class="house-pay-way f16"]//text()')[0]) + '元/月'
            else:
                gfp['价格'] = ' '

        except Exception as ex_results:
            gfp['价格'] = ' '
            continue
            #print(gfp)
        #return gfp
        #return pd.DataFrame(gfp,index=[0])
        getinfo_page.append(pd.DataFrame(gfp,index=[0]))
    return getinfo_page



def Write_info(getinfo_page):

    #City = Pinyin().get_pinyin(u"{}".format(city))
    pd.concat(getinfo_page, ignore_index=True).to_excel('zufang_info.xlsx',index=False)





if __name__ == '__main__':
    t = random.uniform(10,15)
    Base_url = 'https://zk.58.com/chuzu/pn{}/?utm_source=market&spm=b-31580022738699-me-f-862.mingzhan&PGTID=0d3090a7-003a-5338-3ae4-adbbf66a341c&ClickID=2'
    #url = Base_url.format(input('请输入页码'))
    getin_page = []
    for i in range(10):
        if i == 0:
            url = 'https://zk.58.com/chuzu/?utm_source=market&spm=b-31580022738699-me-f-862.mingzhan&PGTID=0d100000-003a-5e86-2ad0-3d263d29e0b9&ClickID=4'
        else:
            url = Base_url.format(i+1)
        gp = get_page(url)
        gu = getinfo_url(gp)
        gf = getinfo(gu)

        getin_page.extend(gf)
        print('第{}页'.format(i + 1))


        time.sleep(t)
    pd.concat(getin_page,ignore_index=True,sort=True).to_excel('zufang.xlsx', index=False)
