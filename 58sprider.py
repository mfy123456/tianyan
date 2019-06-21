from selenium import webdriver
import urllib
from pyquery import PyQuery as pq
import json
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
from selenium.common.exceptions import TimeoutException
import time
from selenium.webdriver import ActionChains
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.wait import WebDriverWait
from PIL import Image

import base64
from io import BytesIO
from fontTools.ttLib import TTFont
import requests
import re
from lxml import etree
import re,random
import pandas as pd
from openpyxl.workbook import Workbook
from xpinyin import Pinyin

# 正则获取字体文件内容
#url = 'https://cd.58.com/wuhou/chuzu/b5j5'

url = 'https://zk.58.com/chuzu/pn3/?utm_source=market&spm=b-31580022738699-me-f-862.mingzhan&PGTID=0d3090a7-003a-5338-3ae4-adbbf66a341c&ClickID=2'
headers = {
    "accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3",
    "accept-encoding": "gzip, deflate, br",
    "accept-language": "zh-CN,zh;q=0.9",
    "cache-control": "max-age=0",
    "upgrade-insecure-requests": "1",
    "user-agent": "Mozilla/5.0 (Windows NT 6.1; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/73.0.3683.86 Safari/537.36",
}


res = requests.get(url, headers=headers)
bs64_str = re.findall("charset=utf-8;base64,(.*?)'\)", res.text)[0]

def get_page_show_ret(string):
    font = TTFont(BytesIO(base64.decodebytes(bs64_str.encode())))
    c = font['cmap'].tables[0].ttFont.tables['cmap'].tables[0].cmap
    ret_list = []
    for char in string:
        decode_num = ord(char)
        if decode_num in c:
            num = c[decode_num]
            num = int(num[-2:])-1
            ret_list.append(num)
        else:
            ret_list.append(char)
    ret_str_show = ''
    for num in ret_list:
        ret_str_show += str(num)
    return ret_str_show

page = etree.HTML(res.text)
li = page.xpath('.//ul[@class="house-list"]//li')[0:-1]

for each_li in li:
    title1 = each_li.xpath('.//div[@class="des"]/h2/a/text()')[0].strip()
    title = get_page_show_ret(title1)
    price1 = each_li.xpath('.//div[@class="money"]/b/text()')[0]
    price = get_page_show_ret(price1)
    info_url = each_li.xpath('.//div[@class="des"]/h2/a')[0].get('href')




